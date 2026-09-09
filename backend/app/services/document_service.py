"""Extraction, chunking, embedding, and durable storage for uploaded documents."""
import asyncio
import mimetypes
import os
import tempfile
import uuid

from fastapi import UploadFile
from app.core.logging import get_logger
from app.services.supabase_service import create_document_metadata, delete_file, update_document_metadata, upload_bytes
from app.vectorstore.embedding_pipeline import ingest_document

logger = get_logger(__name__)
MAX_FILE_SIZE_BYTES = 50 * 1024 * 1024


def _safe_filename(session_id: int, original_name: str) -> str:
    name = os.path.basename(original_name).strip()
    name = "".join(char if char.isalnum() or char in ".-_ " else "_" for char in name)[:200]
    return f"session_{session_id}_{name or 'upload'}"


def _csv_excel_pages(path: str, filename: str) -> list[dict]:
    import pandas as pd
    dataframe = pd.read_csv(path, dtype=str, na_filter=False) if filename.lower().endswith(".csv") else pd.read_excel(path, dtype=str).fillna("")
    if dataframe.empty:
        raise IOError(f"File '{filename}' appears to be empty.")
    pages = [{"text": f"Columns ({len(dataframe.columns)}): {', '.join(dataframe.columns)}\nTotal rows: {len(dataframe)}", "page": 1}]
    for number, start in enumerate(range(0, len(dataframe), 50), start=2):
        rows = []
        for _, row in dataframe.iloc[start:start + 50].iterrows():
            rows.append(" | ".join(f"{column}: {value}" for column, value in row.items() if str(value).strip()))
        pages.append({"text": f"Rows {start + 1}-{start + len(rows)}:\n" + "\n".join(rows), "page": number})
    return pages


def _docx_pages(path: str, filename: str) -> list[dict]:
    from docx import Document
    try:
        document = Document(path)
    except Exception as exc:
        raise IOError(f"Failed to read DOCX '{filename}': {exc}") from exc
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows)
    text = "\n".join(line for line in lines if line.strip()).strip()
    if not text:
        raise IOError(f"No text could be extracted from '{filename}'.")
    return [{"text": text, "page": 1}]


def _txt_pages(content: bytes, filename: str) -> list[dict]:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = content.decode(encoding).strip()
            if text:
                return [{"text": text, "page": 1}]
        except UnicodeDecodeError:
            pass
    raise IOError(f"No readable text could be extracted from '{filename}'.")


def _pdf_pages(path: str, filename: str) -> list[dict]:
    try:
        from langchain_community.document_loaders import PyMuPDFLoader
        loaded = PyMuPDFLoader(path).load()
    except ImportError as exc:
        raise ImportError("PyMuPDF is required for PDF processing. Run: pip install pymupdf") from exc
    except Exception as exc:
        raise IOError(f"Failed to extract text from PDF '{filename}': {exc}") from exc
    pages = [{"text": page.page_content, "page": page.metadata.get("page", 0) + 1} for page in loaded if page.page_content.strip()]
    if not pages:
        raise IOError(f"No text could be extracted from '{filename}'. It may be scanned, encrypted, or empty.")
    return pages


def _index_pages(session_id: int, filename: str, pages: list[dict], document_id: str) -> tuple[int, int]:
    chunks = 0
    characters = 0
    for source in pages:
        text = source["text"].strip()
        if not text:
            continue
        page = source["page"]
        metadata = {"document_id": document_id, "session_id": session_id, "filename": filename, "page": page, "total_pages": len(pages), "source": f"{filename}:p{page}"}
        chunks += ingest_document(session_id, text, metadata, document_id)
        characters += len(text)
    if chunks == 0:
        raise IOError(f"No indexable text could be extracted from '{filename}'.")
    return chunks, characters


async def process_and_store_document(file: UploadFile, session_id: int) -> tuple[int, str]:
    """Store a file, extract text, embed chunks, and write status metadata to Supabase."""
    content = await file.read()
    if not content:
        raise ValueError("The uploaded file is empty.")
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise ValueError("File too large. Maximum size is 50MB.")
    filename = file.filename or "upload"
    ext = os.path.splitext(filename)[1].lower()
    if ext not in {".pdf", ".docx", ".txt", ".csv", ".xls", ".xlsx"}:
        raise ValueError("Unsupported file type. Supported: PDF, DOCX, TXT, CSV, XLS, XLSX.")

    storage_path = f"uploads/session_{session_id}/{_safe_filename(session_id, filename)}"
    content_type = file.content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"
    document_id = str(uuid.uuid4())
    upload_bytes(storage_path, content, content_type)
    try:
        create_document_metadata(document_id, session_id, filename, content_type, storage_path, len(content))
    except Exception as exc:
        delete_file(storage_path)
        raise IOError("Failed to create document metadata in Supabase.") from exc

    temporary_path = ""
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temporary:
            temporary.write(content)
            temporary_path = temporary.name
        if ext in {".csv", ".xls", ".xlsx"}:
            pages = await asyncio.to_thread(_csv_excel_pages, temporary_path, filename)
        elif ext == ".docx":
            pages = await asyncio.to_thread(_docx_pages, temporary_path, filename)
        elif ext == ".txt":
            pages = _txt_pages(content, filename)
        else:
            pages = await asyncio.to_thread(_pdf_pages, temporary_path, filename)
        chunks, characters = await asyncio.to_thread(_index_pages, session_id, filename, pages, document_id)
        update_document_metadata(document_id, status="ready", chunk_count=chunks, extracted_characters=characters, error=None)
        logger.info("Document indexed", extra={"session_id": session_id, "document_id": document_id, "chunks": chunks})
        return chunks, storage_path
    except (ValueError, ImportError, IOError) as exc:
        update_document_metadata(document_id, status="failed", error=str(exc)[:1000])
        raise
    except Exception as exc:
        logger.exception("Document processing failed", extra={"session_id": session_id})
        update_document_metadata(document_id, status="failed", error=str(exc)[:1000])
        raise IOError(f"Failed to process document: {exc}") from exc
    finally:
        if temporary_path and os.path.exists(temporary_path):
            os.remove(temporary_path)
